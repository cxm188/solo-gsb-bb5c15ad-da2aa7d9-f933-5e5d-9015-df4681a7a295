"""FastAPI 入口：上传 XLSX、构建公式依赖图、提供影响分析与导出。"""

from __future__ import annotations

import io
import os
import re
import uuid
from dataclasses import dataclass
from typing import Dict, List, Optional

from fastapi import FastAPI, File, HTTPException, Query, UploadFile
from fastapi.responses import FileResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string

from .analyzer import (
    Analysis,
    _coord_sort_key,
    bfs_shortest,
    enumerate_paths,
)

MAX_UPLOAD_BYTES = int(os.environ.get("MAX_UPLOAD_MB", "20")) * 1024 * 1024
# 单次影响分析最多返回的受影响单元格数量
MAX_AFFECTED = int(os.environ.get("MAX_AFFECTED", "20000"))

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
STATIC_DIR = os.path.join(BASE_DIR, "static")

app = FastAPI(title="公式影响分析", version="1.0.0")


@dataclass
class Stored:
    filename: str
    analysis: Analysis


STORE: Dict[str, Stored] = {}


# --------------------------------------------------------------------------- #
# 辅助
# --------------------------------------------------------------------------- #
def _normalize_coord(coord: str) -> Optional[str]:
    coord = coord.strip().replace("$", "").upper()
    m = re.fullmatch(r"[A-Z]{1,3}[0-9]{1,7}", coord)
    if not m:
        return None
    try:
        col_letters, row = coordinate_from_string(coord)
        col = column_index_from_string(col_letters)
    except Exception:
        return None
    if col > 16384 or row > 1048576:
        return None
    return coord


def _get(wb_id: str) -> Stored:
    stored = STORE.get(wb_id)
    if stored is None:
        raise HTTPException(status_code=404, detail="工作簿不存在或已过期，请重新上传")
    return stored


def _node_payload(analysis: Analysis, node_id: int) -> Dict:
    n = analysis.nodes[node_id]
    return {
        "id": n.id,
        "key": n.key,
        "sheet": n.sheet,
        "coord": n.coord,
        "value": n.value,
        "formula": n.formula,
    }


def _path_payload(analysis: Analysis, path_ids: List[int]) -> List[Dict]:
    return [
        {
            "key": analysis.nodes[i].key,
            "sheet": analysis.nodes[i].sheet,
            "coord": analysis.nodes[i].coord,
            "formula": analysis.nodes[i].formula,
        }
        for i in path_ids
    ]


def _sheet_order(analysis: Analysis) -> Dict[str, int]:
    return {s["name"]: i for i, s in enumerate(analysis.sheets)}


# --------------------------------------------------------------------------- #
# 接口
# --------------------------------------------------------------------------- #
@app.post("/api/workbooks")
async def upload_workbook(file: UploadFile = File(...)):
    filename = file.filename or "uploaded.xlsx"
    if not filename.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(status_code=400, detail="仅支持 .xlsx / .xlsm 文件")

    data = await file.read()
    if len(data) > MAX_UPLOAD_BYTES:
        raise HTTPException(
            status_code=413,
            detail=f"文件过大（上限 {MAX_UPLOAD_BYTES // 1024 // 1024} MB）",
        )
    if not data:
        raise HTTPException(status_code=400, detail="文件为空")

    try:
        wb = load_workbook(io.BytesIO(data), data_only=False, read_only=False)
    except Exception as exc:  # openpyxl 对坏文件会抛 InvalidFileError / zip 错误
        raise HTTPException(status_code=400, detail=f"无法解析工作簿: {exc}") from exc

    from .analyzer import _Builder

    try:
        analysis = _Builder(wb, filename=filename).build()
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"分析失败: {exc}") from exc

    wb_id = uuid.uuid4().hex[:12]
    STORE[wb_id] = Stored(filename=filename, analysis=analysis)
    return {
        "id": wb_id,
        "filename": filename,
        "summary": analysis.to_dict(),
    }


@app.get("/api/workbooks/{wb_id}")
def get_analysis(wb_id: str):
    return _get(wb_id).analysis.to_dict()


@app.get("/api/workbooks/{wb_id}/nodes")
def list_nodes(
    wb_id: str,
    sheet: Optional[str] = None,
    formula_only: bool = False,
    q: Optional[str] = Query(None, description="按坐标或公式文本过滤"),
):
    analysis = _get(wb_id).analysis
    order = _sheet_order(analysis)
    rows = []
    for n in analysis.nodes:
        if sheet is not None and sheet != "" and n.sheet != sheet:
            continue
        if formula_only and n.formula is None:
            continue
        if q:
            needle = q.strip().upper()
            hay = f"{n.coord} {n.formula or ''} {n.value or ''}".upper()
            if needle not in hay:
                continue
        rows.append(
            {
                "id": n.id,
                "key": n.key,
                "sheet": n.sheet,
                "sheet_order": order.get(n.sheet, 999),
                "coord": n.coord,
                "formula": n.formula,
                "value": n.value,
            }
        )
    rows.sort(key=lambda r: (r["sheet_order"], _coord_sort_key(r["coord"])))
    truncated = len(rows) > 20000
    return {"count": len(rows), "truncated": truncated, "nodes": rows[:20000]}


@app.get("/api/workbooks/{wb_id}/impact")
def impact(
    wb_id: str,
    sheet: str = Query(..., description="工作表名"),
    coord: str = Query(..., description="单元格坐标，如 B2"),
    direction: str = Query("dependents", pattern="^(dependents|precedents)$"),
):
    """影响分析。

    direction=dependents：沿引用反向追溯，列出所有（直接/间接）受所选单元格影响的公式格；
    direction=precedents：正向列出所选单元格依赖的所有上游单元格。
    """
    analysis = _get(wb_id).analysis
    norm = _normalize_coord(coord)
    if norm is None:
        raise HTTPException(status_code=400, detail=f"非法单元格坐标: {coord}")

    node_id = analysis._key_to_id.get(f"{sheet}!{norm}")
    if node_id is None:
        raise HTTPException(
            status_code=404,
            detail=f"工作表 {sheet} 中不存在参与公式依赖的单元格 {norm}",
        )

    adj = analysis._in if direction == "dependents" else analysis._out
    shortest = bfs_shortest(node_id, adj)

    cycle_keys = {k for cyc in analysis.cycles for k in cyc.nodes}
    order = _sheet_order(analysis)

    affected_ids = [i for i in shortest if i != node_id]
    truncated = len(affected_ids) > MAX_AFFECTED
    affected_ids.sort(
        key=lambda i: (
            len(shortest[i]),
            order.get(analysis.nodes[i].sheet, 999),
            _coord_sort_key(analysis.nodes[i].coord),
        )
    )
    affected_ids = affected_ids[:MAX_AFFECTED]

    entries = []
    direct_ids: List[int] = []
    for tid in affected_ids:
        dist = len(shortest[tid]) - 1
        sp = shortest[tid]
        paths = _path_payload(analysis, sp)
        # 图较小时多给几条不同的传播路径，帮助排查
        if dist > 1 and len(shortest) <= 3000:
            alt = enumerate_paths(node_id, tid, adj, max_paths=3)
            paths_multi = [_path_payload(analysis, p) for p in alt] or [paths]
        else:
            paths_multi = [paths]
        node = analysis.nodes[tid]
        if dist == 1:
            direct_ids.append(tid)
        entries.append(
            {
                "node": _node_payload(analysis, tid),
                "distance": dist,
                "relation": "direct" if dist == 1 else "indirect",
                "on_cycle": node.key in cycle_keys,
                "paths": paths_multi,
            }
        )

    direct = sorted(
        direct_ids,
        key=lambda i: (
            order.get(analysis.nodes[i].sheet, 999),
            _coord_sort_key(analysis.nodes[i].coord),
        ),
    )
    return {
        "workbook_id": wb_id,
        "direction": direction,
        "selected": _node_payload(analysis, node_id),
        "direct_count": len(direct),
        "indirect_count": sum(1 for e in entries if e["relation"] == "indirect"),
        "affected_count": len(entries),
        "truncated": truncated,
        "affected": entries,
    }


@app.get("/api/workbooks/{wb_id}/export")
def export_analysis(wb_id: str):
    stored = _get(wb_id)
    payload = JSONResponse(stored.analysis.to_dict())
    # 附带下载文件名
    payload.headers["Content-Disposition"] = (
        f'attachment; filename="analysis-{wb_id}.json"'
    )
    return payload


@app.get("/api/workbooks")
def list_workbooks():
    return [
        {"id": wid, "filename": s.filename, "stats": s.analysis.stats}
        for wid, s in STORE.items()
    ]


@app.delete("/api/workbooks/{wb_id}")
def delete_workbook(wb_id: str):
    STORE.pop(wb_id, None)
    return {"ok": True}


# --------------------------------------------------------------------------- #
# 静态页面（放在最后，避免吞掉 /api）
# --------------------------------------------------------------------------- #
app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")


@app.get("/")
def index():
    return FileResponse(os.path.join(STATIC_DIR, "index.html"))
