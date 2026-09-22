"""analyzer 的单元测试（标准库 unittest，无额外依赖）。

运行：python -m unittest discover -s tests -v
"""
import os
import unittest

from openpyxl import Workbook

from app.analyzer import (
    MAX_EXPANSION,
    _Builder,
    _norm_coord,
    bfs_shortest,
    enumerate_paths,
)

SAMPLE = os.path.join(
    os.path.dirname(__file__), "..", "examples", "sample_formulas.xlsx"
)


def build(wb, name="test.xlsx"):
    return _Builder(wb, name).build()


def edges_by_formula(analysis):
    return analysis.edges


class CoordTests(unittest.TestCase):
    def test_norm(self):
        self.assertEqual(_norm_coord("a1"), "A1")
        self.assertEqual(_norm_coord("$XFD$1048576"), "XFD1048576")
        self.assertIsNone(_norm_coord("XFE1"))      # 超出 16384 列
        self.assertIsNone(_norm_coord("A1048577"))  # 超出行
        self.assertIsNone(_norm_coord("LOG"))       # 仅字母
        self.assertIsNone(_norm_coord("12"))        # 仅数字


class BasicGraphTests(unittest.TestCase):
    def test_simple_chain_and_cross_sheet(self):
        wb = Workbook()
        ws1 = wb.active
        ws1.title = "S1"
        ws1["A1"] = 1
        ws1["A2"] = "=A1*2"
        ws1["A3"] = "=A2+A1"
        ws2 = wb.create_sheet("S2")
        ws2["B1"] = "=S1!A3"

        a = build(wb)
        self.assertEqual(a.edges["S1!A2"], ["S1!A1"])
        self.assertEqual(set(a.edges["S1!A3"]), {"S1!A1", "S1!A2"})
        self.assertEqual(a.edges["S2!B1"], ["S1!A3"])
        self.assertEqual(a.stats["cross_sheet_edge_count"], 1)
        self.assertEqual(a.stats["cycle_count"], 0)
        # 常量值被记录
        node = next(n for n in a.nodes if n.key == "S1!A1")
        self.assertEqual(node.value, "1")
        self.assertIsNone(node.formula)

    def test_rectangle_expansion(self):
        wb = Workbook()
        ws = wb.active
        ws["D5"] = "=SUM(A1:B2)"
        a = build(wb)
        targets = set(a.edges["Sheet!D5"])
        self.assertEqual(
            targets, {"Sheet!A1", "Sheet!A2", "Sheet!B1", "Sheet!B2"}
        )

    def test_quoted_sheet_name(self):
        wb = Workbook()
        ws1 = wb.active
        ws1.title = "my data"
        ws1["A1"] = 10
        ws2 = wb.create_sheet("other")
        ws2["C3"] = "='my data'!A1"
        a = build(wb)
        self.assertEqual(a.edges["other!C3"], ["my data!A1"])

    def test_dollar_abs_refs(self):
        wb = Workbook()
        ws = wb.active
        ws["B2"] = "=$A$1+A$2"
        a = build(wb)
        self.assertEqual(set(a.edges["Sheet!B2"]), {"Sheet!A1", "Sheet!A2"})

    def test_external_ref(self):
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "=[1]Sheet1!B2*2"
        a = build(wb)
        self.assertNotIn("Sheet!A1", a.edges)
        self.assertIn("[1]Sheet1!B2", a.external_refs)

    def test_missing_sheet_warns(self):
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "=Ghost!B9"
        a = build(wb)
        # 仍然建节点（文件可能只是缺页）
        self.assertEqual(a.edges["Sheet!A1"], ["Ghost!B9"])
        self.assertTrue(any("不存在" in w for w in a.warnings))


class CycleTests(unittest.TestCase):
    def test_three_node_cycle(self):
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "=B1+1"
        ws["B1"] = "=C1+1"
        ws["C1"] = "=A1+1"
        a = build(wb)
        self.assertEqual(a.stats["cycle_count"], 1)
        cyc = a.cycles[0].nodes
        self.assertEqual(cyc[0], cyc[-1])
        self.assertEqual(set(cyc), {"Sheet!A1", "Sheet!B1", "Sheet!C1"})

    def test_self_loop(self):
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "=A1+1"
        a = build(wb)
        self.assertEqual(a.stats["cycle_count"], 1)
        self.assertEqual(a.cycles[0].nodes, ["Sheet!A1", "Sheet!A1"])
        self.assertEqual(a.edges["Sheet!A1"], ["Sheet!A1"])

    def test_diamond_is_not_cycle(self):
        # A -> B,C; B,C -> D
        wb = Workbook()
        ws = wb.active
        ws["B1"] = "=A1"
        ws["C1"] = "=A1"
        ws["D1"] = "=B1+C1"
        a = build(wb)
        self.assertEqual(a.stats["cycle_count"], 0)


class ImpactTests(unittest.TestCase):
    def _diamond(self):
        wb = Workbook()
        ws = wb.active
        ws["B1"] = "=A1"
        ws["C1"] = "=A1"
        ws["D1"] = "=B1+C1"
        return build(wb)

    def test_dependents_bfs(self):
        a = self._diamond()
        sp = bfs_shortest(a._key_to_id["Sheet!A1"], a._in)
        distances = {a.nodes[i].coord: len(p) - 1 for i, p in sp.items()}
        self.assertEqual(distances["B1"], 1)
        self.assertEqual(distances["C1"], 1)
        self.assertEqual(distances["D1"], 2)

    def test_precedents(self):
        a = self._diamond()
        sp = bfs_shortest(a._key_to_id["Sheet!D1"], a._out)
        reached = {a.nodes[i].coord for i in sp}
        self.assertEqual(reached, {"D1", "B1", "C1", "A1"})

    def test_multiple_paths(self):
        a = self._diamond()
        paths = enumerate_paths(
            a._key_to_id["Sheet!A1"], a._key_to_id["Sheet!D1"], a._in
        )
        self.assertEqual(len(paths), 2)
        chains = [
            tuple(a.nodes[i].coord for i in p)
            for p in paths
        ]
        self.assertIn(("A1", "B1", "D1"), chains)
        self.assertIn(("A1", "C1", "D1"), chains)


class DefinedNameTests(unittest.TestCase):
    def test_workbook_defined_name_range(self):
        from openpyxl.workbook.defined_name import DefinedName

        wb = Workbook()
        ws = wb.active
        for i in range(1, 4):
            ws.cell(row=i, column=1, value=i)
        wb.defined_names["数据区"] = DefinedName(
            name="数据区", attr_text="Sheet!$A$1:$A$3"
        )
        ws["B1"] = "=SUM(数据区)"
        a = build(wb)
        self.assertEqual(
            set(a.edges["Sheet!B1"]),
            {"Sheet!A1", "Sheet!A2", "Sheet!A3"},
        )


class FullColumnTests(unittest.TestCase):
    def test_full_column_limited_by_dimension(self):
        wb = Workbook()
        ws = wb.active
        ws["A1"] = 1
        ws["A2"] = 2
        ws["C1"] = "=SUM(A:A)"
        a = build(wb)
        targets = a.edges["Sheet!C1"]
        # 维度是 A1:C2 左右 -> 展开 2 行，远小于上限
        self.assertEqual(targets, ["Sheet!A1", "Sheet!A2"])
        self.assertLessEqual(len(targets), MAX_EXPANSION)


class SampleWorkbookTests(unittest.TestCase):
    def setUp(self):
        if not os.path.exists(SAMPLE):
            self.skipTest("示例文件尚未生成: python examples/make_sample.py")

    def test_sample_structure(self):
        from openpyxl import load_workbook

        wb = load_workbook(SAMPLE, data_only=False)
        a = _Builder(wb, "sample_formulas.xlsx").build()
        self.assertGreaterEqual(a.stats["cross_sheet_edge_count"], 4)
        self.assertEqual(a.stats["cycle_count"], 2)
        # B2 的影响可一路传到管理层报告!B4
        sp = bfs_shortest(a._key_to_id["基础数据!B2"], a._in)
        self.assertIn(a._key_to_id["管理层报告!B4"], sp)
        # 定义名称参与计算
        self.assertIn("税率", a.defined_names)


if __name__ == "__main__":
    unittest.main()
